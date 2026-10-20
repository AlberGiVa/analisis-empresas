"""Excel workbook import: template, round trip, error messages and tolerant parsing."""

from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from company_analysis.models import DataBundle, FinancialItem
from company_analysis.providers import SampleProvider
from company_analysis.providers.workbook import (
    SHEETS,
    WorkbookError,
    parse_number,
    read_workbook,
    sheet_columns,
    write_workbook,
)
from company_analysis.settings import PROJECT_ROOT


@pytest.fixture(scope="module")
def sample() -> DataBundle:
    return SampleProvider(PROJECT_ROOT / "data" / "sample").fetch()


@pytest.fixture
def example(sample: DataBundle, tmp_path: Path) -> Path:
    return write_workbook(sample, tmp_path / "ejemplo.xlsx")


def edit(path: Path, sheet: str, row: int, header: str, value: str | float | None) -> None:
    """Change (or empty) one cell, found by its column header, as an analyst would in Excel."""
    workbook = load_workbook(path)
    worksheet = workbook[sheet]
    assert isinstance(worksheet, Worksheet)
    headers = [cell.value for cell in worksheet[1]]
    # Assign .value: openpyxl's cell(..., value=None) would leave the old value in place.
    worksheet.cell(row=row, column=headers.index(header) + 1).value = value
    workbook.save(path)


def test_the_sample_survives_a_round_trip_through_excel(sample: DataBundle, example: Path) -> None:
    imported = read_workbook(example)
    for collection in sample.row_counts():
        assert set(getattr(imported.bundle, collection)) == set(getattr(sample, collection))
    assert imported.warnings == ()


def test_template_has_instructions_dropdowns_and_every_sheet(tmp_path: Path) -> None:
    path = write_workbook(DataBundle(), tmp_path / "plantilla.xlsx")
    workbook = load_workbook(path)
    assert workbook.sheetnames[0] == "Instrucciones"
    assert workbook["Listas"].sheet_state == "hidden"
    for sheet in SHEETS:
        worksheet = workbook[sheet.name]
        assert isinstance(worksheet, Worksheet)
        assert [cell.value for cell in worksheet[1]] == [c.header for c in sheet_columns(sheet)]
    issuers = workbook["Emisores"]
    assert isinstance(issuers, Worksheet)
    assert len(issuers.data_validations.dataValidation) == 4  # país, sector, industria, divisa


def test_an_empty_template_cannot_be_loaded(tmp_path: Path) -> None:
    path = write_workbook(DataBundle(), tmp_path / "plantilla.xlsx")
    with pytest.raises(WorkbookError, match="La hoja «Emisores» está vacía"):
        read_workbook(path)


def test_errors_point_to_sheet_row_and_column(example: Path) -> None:
    edit(example, "Emisores", 2, "Sector", "Inventado")
    edit(example, "Bonos", 2, "Vencimiento", "31/02/2026")
    edit(example, "Acciones", 3, "ISIN", "XX123")
    edit(example, "Estados financieros", 2, "Ventas", "mucho")
    with pytest.raises(WorkbookError) as raised:
        read_workbook(example)
    errors = raised.value.errors
    assert any(
        e.startswith("Hoja «Emisores», fila 2, columna «Sector»: valor no admitido") for e in errors
    )
    assert any("Hoja «Bonos», fila 2, columna «Vencimiento»: fecha no válida" in e for e in errors)
    assert any("Hoja «Acciones», fila 3, columna «ISIN»: ISIN no válido" in e for e in errors)
    assert any("columna «Ventas»: se esperaba un número" in e for e in errors)


def test_inconsistent_references_are_reported(example: Path) -> None:
    edit(example, "Acciones", 2, "Emisor", "NOEXISTE")
    with pytest.raises(WorkbookError, match="emisor desconocido: NOEXISTE"):
        read_workbook(example)


def test_numbers_and_dates_typed_as_text_in_spanish(example: Path) -> None:
    edit(example, "Cotizaciones", 2, "Cierre", "1.234,5")
    edit(example, "Cotizaciones", 2, "Fecha", "06/01/2023")
    first_price = read_workbook(example).bundle.equity_prices[0]
    assert first_price.close == 1234.5
    assert first_price.price_date == date(2023, 1, 6)


def test_empty_cells_stay_missing(example: Path) -> None:
    edit(example, "Estados financieros", 2, "Inversiones (capex)", None)
    facts = read_workbook(example).bundle.financial_facts
    first = facts[0]
    same_period = [
        f for f in facts if (f.issuer_id, f.period_end) == (first.issuer_id, first.period_end)
    ]
    assert FinancialItem.CAPITAL_EXPENDITURE not in {f.item for f in same_period}


def test_publication_on_the_closing_date_is_flagged(example: Path) -> None:
    workbook = load_workbook(example)
    worksheet = workbook["Estados financieros"]
    assert isinstance(worksheet, Worksheet)
    closing = worksheet.cell(row=2, column=3).value  # "Fecha de cierre"
    worksheet.cell(row=2, column=4, value=closing)  # "Fecha de publicación"
    workbook.save(example)
    (warning,) = read_workbook(example).warnings
    assert warning.startswith("Hoja «Estados financieros», fila 2: la fecha de publicación")


@pytest.mark.parametrize(
    ("text", "number"),
    [
        ("1.234,5", 1234.5),
        ("1,234.5", 1234.5),
        ("12,5", 12.5),
        ("12.5", 12.5),
        ("1.234.567", 1234567.0),
        ("-3,2", -3.2),
        (f"1{chr(0xA0)}234,5", 1234.5),  # non-breaking space as thousands separator
    ],
)
def test_parse_number(text: str, number: float) -> None:
    assert parse_number(text) == pytest.approx(number)
