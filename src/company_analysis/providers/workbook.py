"""Import of real data from an Excel workbook that follows the fund's template (in Spanish).

Analysts fill the template with the Bloomberg or FactSet Excel add-ins (or by hand) and the tool
validates it and loads it. One workbook is one load: it must contain the issuers and instruments
its data refer to. Empty cells mean "no data" and are never filled in.
"""

from collections import defaultdict
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import ClassVar, Literal

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import BaseModel, ValidationError

from company_analysis.models import (
    Bond,
    BondQuote,
    DataBundle,
    Equity,
    EquityPrice,
    FinancialFact,
    FinancialItem,
    FiscalPeriod,
    FxRate,
    Industry,
    Issuer,
    RatingAction,
    Sector,
    Seniority,
    SpreadType,
)
from company_analysis.models.identifiers import SUPPORTED_COUNTRIES, SUPPORTED_CURRENCIES
from company_analysis.providers.base import DataProvider

Kind = Literal["text", "number", "integer", "date"]
MAX_REPORTED_ERRORS = 40
LISTS_SHEET = "Listas"
LAST_VALIDATED_ROW = 5000


@dataclass(frozen=True)
class Choices:
    """Values allowed in a column: label shown in the workbook -> stored value."""

    name: str  # title of the list in the hidden "Listas" sheet
    values: Mapping[str, str]


@dataclass(frozen=True)
class Column:
    header: str  # shown in the workbook (Spanish)
    field: str  # model field it fills
    kind: Kind = "text"
    required: bool = True
    choices: Choices | None = None
    help: str = ""
    bloomberg: str = "—"  # orientative source fields; to be checked in each terminal
    factset: str = "—"


@dataclass(frozen=True)
class Sheet:
    name: str
    collection: str  # DataBundle collection
    model: type[BaseModel]
    description: str
    columns: tuple[Column, ...]
    required: bool = False


COUNTRIES = Choices("Países", {code: code for code in sorted(SUPPORTED_COUNTRIES)})
SECTORS = Choices("Sectores", {sector.label: sector.value for sector in Sector})
INDUSTRIES = Choices("Industrias", {industry.label: industry.value for industry in Industry})
CURRENCIES = Choices("Divisas", {code: code for code in sorted(SUPPORTED_CURRENCIES)})
FX_CURRENCIES = Choices(
    "Divisas frente al euro", {code: code for code in sorted(SUPPORTED_CURRENCIES - {"EUR"})}
)
SENIORITIES = Choices("Prelaciones", {seniority.label: seniority.value for seniority in Seniority})
PERIODS = Choices("Periodos", {period.value: period.value for period in FiscalPeriod})
SPREAD_TYPES = Choices(
    "Tipos de diferencial",
    {
        "Z-spread": SpreadType.Z_SPREAD.value,
        "G-spread": SpreadType.G_SPREAD.value,
        "OAS": SpreadType.OAS.value,
    },
)
ALL_CHOICES = (
    COUNTRIES, SECTORS, INDUSTRIES, CURRENCIES, FX_CURRENCIES, SENIORITIES, PERIODS, SPREAD_TYPES,
)  # fmt: skip

ISSUER_REFERENCE = Column("Emisor", "issuer_id", help="Identificador de la hoja Emisores.")

SHEETS: tuple[Sheet, ...] = (
    Sheet(
        "Emisores",
        "issuers",
        Issuer,
        "Una fila por empresa.",
        (
            Column(
                "Identificador",
                "issuer_id",
                help="Código corto y único que eliges (2 a 12 letras o números en mayúsculas), "
                "por ejemplo el ticker sin el mercado.",
            ),
            Column("Nombre", "name", help="Denominación social.", bloomberg="NAME"),
            Column(
                "LEI",
                "lei",
                required=False,
                help="Identificador LEI del emisor (20 caracteres).",
                bloomberg="LEGAL_ENTITY_IDENTIFIER",
            ),
            Column(
                "País",
                "country",
                choices=COUNTRIES,
                help="País del emisor (código ISO de dos letras).",
                bloomberg="CNTRY_OF_DOMICILE",
            ),
            Column(
                "Sector",
                "sector",
                choices=SECTORS,
                help="Clasificación propia del fondo (lista desplegable).",
                bloomberg="Asignar a partir de BICS_LEVEL_1_SECTOR_NAME",
                factset="Asignar a partir de su clasificación sectorial",
            ),
            Column(
                "Industria",
                "industry",
                choices=INDUSTRIES,
                help="Debe pertenecer al sector elegido. Bancos y aseguradoras determinan qué "
                "métricas no aplican.",
            ),
            Column(
                "Divisa de reporte",
                "reporting_currency",
                choices=CURRENCIES,
                help="Divisa en la que presenta sus estados financieros.",
                bloomberg="EQY_FUND_CRNCY",
            ),
            Column(
                "Mes de cierre fiscal",
                "fiscal_year_end_month",
                "integer",
                help="Número del mes en que cierra su ejercicio (12 = diciembre).",
                bloomberg="EQY_FISCAL_YR_END (mes)",
            ),
        ),
        required=True,
    ),
    Sheet(
        "Acciones",
        "equities",
        Equity,
        "La acción cotizada de cada empresa (exactamente una por empresa).",
        (
            Column("ISIN", "isin", help="ISIN de la acción.", bloomberg="ID_ISIN"),
            ISSUER_REFERENCE,
            Column("Ticker", "ticker", help="Ticker de la acción.", bloomberg="TICKER"),
            Column(
                "Mercado (MIC)",
                "exchange",
                help="Código MIC de 4 letras del mercado principal (por ejemplo, XMAD).",
                bloomberg="ID_MIC_PRIM_EXCH",
            ),
            Column(
                "Divisa de cotización",
                "currency",
                choices=CURRENCIES,
                help="Divisa del precio (las acciones británicas, en libras y no en peniques).",
                bloomberg="CRNCY",
            ),
        ),
        required=True,
    ),
    Sheet(
        "Bonos",
        "bonds",
        Bond,
        "Bonos de cada empresa. Una empresa sin bonos no aparece en esta hoja.",
        (
            Column("ISIN", "isin", help="ISIN del bono.", bloomberg="ID_ISIN"),
            ISSUER_REFERENCE,
            Column("Divisa", "currency", choices=CURRENCIES, bloomberg="CRNCY"),
            Column("Cupón (%)", "coupon_pct", "number", help="Por ejemplo, 4,25.", bloomberg="CPN"),
            Column(
                "Pagos de cupón al año",
                "coupon_frequency",
                "integer",
                help="1 = anual, 2 = semestral, 4 = trimestral.",
                bloomberg="CPN_FREQ",
            ),
            Column("Fecha de emisión", "issue_date", "date", bloomberg="ISSUE_DT"),
            Column("Vencimiento", "maturity_date", "date", bloomberg="MATURITY"),
            Column(
                "Prelación",
                "seniority",
                choices=SENIORITIES,
                help="Orden de cobro en caso de impago (lista desplegable).",
                bloomberg="Traducir PAYMENT_RANK a la lista",
            ),
            Column(
                "Nominal vivo (millones)",
                "amount_outstanding_mn",
                "number",
                help="En millones de la divisa del bono.",
                bloomberg="AMT_OUTSTANDING / 1.000.000",
            ),
        ),
    ),
    Sheet(
        "Estados financieros",
        "financial_facts",
        FinancialFact,
        "Una fila por empresa, periodo y fecha de publicación, y una columna por partida. Importes "
        "en millones de la divisa de reporte. Deja vacía la celda de una partida sin dato: nunca "
        "pongas cero si no lo es.",
        (
            ISSUER_REFERENCE,
            Column(
                "Periodo",
                "fiscal_period",
                choices=PERIODS,
                help="FY = ejercicio completo; Q1, H1 y 9M = acumulado desde el inicio del "
                "ejercicio.",
            ),
            Column("Fecha de cierre", "period_end", "date", help="Último día del periodo."),
            Column(
                "Fecha de publicación",
                "published_at",
                "date",
                help="Día en que la empresa publicó estas cifras, no la fecha de cierre: así nunca "
                "se usan datos antes de conocerse.",
                bloomberg="ANNOUNCEMENT_DT",
            ),
            Column(
                "Divisa",
                "currency",
                choices=CURRENCIES,
                help="Divisa de reporte (la misma que en la hoja Emisores).",
                bloomberg="EQY_FUND_CRNCY",
            ),
        ),
    ),
    Sheet(
        "Cotizaciones",
        "equity_prices",
        EquityPrice,
        "Precios de cierre de las acciones: una fila por acción y fecha.",
        (
            Column("ISIN", "isin", help="ISIN de la acción (hoja Acciones)."),
            Column("Fecha", "price_date", "date"),
            Column("Cierre", "close", "number", bloomberg="PX_LAST", factset="P_PRICE"),
            Column("Divisa", "currency", choices=CURRENCIES, bloomberg="CRNCY"),
        ),
    ),
    Sheet(
        "Cotizaciones de bonos",
        "bond_quotes",
        BondQuote,
        "Rentabilidad y diferencial de los bonos: una fila por bono y fecha.",
        (
            Column("ISIN", "isin", help="ISIN del bono (hoja Bonos)."),
            Column("Fecha", "quote_date", "date"),
            Column("Rentabilidad (%)", "yield_pct", "number", bloomberg="YLD_YTM_MID"),
            Column(
                "Diferencial (pb)",
                "spread_bp",
                "number",
                help="En puntos básicos.",
                bloomberg="Z_SPRD_MID",
            ),
            Column(
                "Tipo de diferencial",
                "spread_type",
                choices=SPREAD_TYPES,
                help="Solo se comparan bonos con el mismo tipo de diferencial.",
            ),
        ),
    ),
    Sheet(
        "Ratings",
        "ratings",
        RatingAction,
        "Historial de ratings de los bonos: una fila por cada cambio de rating de cada agencia.",
        (
            Column("ISIN", "isin", help="ISIN del bono (hoja Bonos)."),
            Column("Agencia", "agency", help="Por ejemplo, S&P, Moody's o Fitch."),
            Column(
                "Rating",
                "rating",
                help="Escala de S&P/Fitch (BBB-) o de Moody's (Baa3).",
                bloomberg="RTG_SP, RTG_MOODY, RTG_FITCH",
            ),
            Column("Fecha", "rating_date", "date", help="Fecha de la acción de rating."),
        ),
    ),
    Sheet(
        "Tipos de cambio",
        "fx_rates",
        FxRate,
        "Tipos de cambio diarios frente al euro, necesarios para cualquier divisa distinta del "
        "euro.",
        (
            Column("Divisa", "currency", choices=FX_CURRENCIES),
            Column("Fecha", "rate_date", "date"),
            Column(
                "Unidades por 1 EUR",
                "units_per_eur",
                "number",
                help="Por ejemplo, 1,08 dólares por euro.",
                bloomberg="PX_LAST de EURUSD Curncy (EURGBP, EURCHF…)",
            ),
        ),
    ),
)

_ITEM_SOURCES: dict[FinancialItem, tuple[str, str]] = {
    FinancialItem.REVENUE: ("SALES_REV_TURN", "FF_SALES"),
    FinancialItem.GROSS_PROFIT: ("GROSS_PROFIT", "FF_GROSS_INC"),
    FinancialItem.EBITDA: ("EBITDA", "FF_EBITDA_OPER"),
    FinancialItem.EBIT: ("EBIT", "FF_EBIT_OPER"),
    FinancialItem.INTEREST_EXPENSE: ("IS_INT_EXPENSE", "FF_INT_EXP_DEBT"),
    FinancialItem.NET_INCOME: ("NET_INCOME", "FF_NET_INC"),
    FinancialItem.OPERATING_CASH_FLOW: ("CF_CASH_FROM_OPER", "FF_OPER_CF"),
    FinancialItem.CAPITAL_EXPENDITURE: ("CAPITAL_EXPEND, con el signo cambiado", "FF_CAPEX"),
    FinancialItem.TOTAL_DEBT: ("SHORT_AND_LONG_TERM_DEBT", "FF_DEBT"),
    FinancialItem.CASH: ("BS_CASH_NEAR_CASH_ITEM", "FF_CASH_ST"),
    FinancialItem.TOTAL_EQUITY: ("TOTAL_EQUITY", "FF_SHLDRS_EQ más minoritarios"),
    FinancialItem.EQUITY_ATTRIBUTABLE: ("TOT_COMMON_EQY", "FF_COM_EQ"),
    FinancialItem.SHARES_OUTSTANDING: ("BS_SH_OUT", "FF_COM_SHS_OUT"),
}
_POSITIVE_ITEMS = {FinancialItem.INTEREST_EXPENSE, FinancialItem.CAPITAL_EXPENDITURE}


def _item_help(item: FinancialItem) -> str:
    if item is FinancialItem.SHARES_OUTSTANDING:
        return "Millones de acciones al cierre del periodo."
    unit = "Millones de la divisa de reporte"
    return f"{unit}, en positivo." if item in _POSITIVE_ITEMS else f"{unit}."


ITEM_COLUMNS = tuple(
    Column(
        item.label,
        item.value,
        "number",
        required=False,
        help=_item_help(item),
        bloomberg=_ITEM_SOURCES[item][0],
        factset=_ITEM_SOURCES[item][1],
    )
    for item in FinancialItem
)
FINANCIALS = next(sheet for sheet in SHEETS if sheet.collection == "financial_facts")


def sheet_columns(sheet: Sheet) -> tuple[Column, ...]:
    """Columns as they appear in the workbook (the financial sheet adds one per line item)."""
    return (*sheet.columns, *ITEM_COLUMNS) if sheet is FINANCIALS else sheet.columns


class WorkbookError(Exception):
    """The workbook cannot be loaded; ``errors`` says where and why, in plain Spanish."""

    def __init__(self, errors: Sequence[str]) -> None:
        super().__init__("\n".join(errors))
        self.errors = list(errors)


@dataclass(frozen=True)
class WorkbookImport:
    bundle: DataBundle
    warnings: tuple[str, ...]


# --- Reading ---------------------------------------------------------------------------------
def parse_number(text: str) -> float:
    """Number typed as text, in Spanish or English notation: "1.234,5", "1,234.5", "12,5"."""
    # str.split() also removes non-breaking spaces (thousands separators in Spanish Excel).
    cleaned = "".join(text.split())
    if "," in cleaned and "." in cleaned:
        decimal = "," if cleaned.rfind(",") > cleaned.rfind(".") else "."
        thousands = "." if decimal == "," else ","
        cleaned = cleaned.replace(thousands, "").replace(decimal, ".")
    elif "," in cleaned:
        cleaned = cleaned.replace(",", ".")
    elif cleaned.count(".") > 1:
        cleaned = cleaned.replace(".", "")
    return float(cleaned)


def _parse_date(text: str) -> date:
    for pattern in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(text.strip(), pattern).date()  # noqa: DTZ007 (only the date)
        except ValueError:
            continue
    raise ValueError(f"fecha no válida «{text}» (usa dd/mm/aaaa)")


def _convert(raw: object, column: Column) -> object:
    """Value of a cell for a model field. Raises ``ValueError`` with a Spanish message."""
    if raw is None or (isinstance(raw, str) and not raw.strip()):
        return None
    if column.choices is not None:
        text = str(raw).strip().casefold()
        for label, value in column.choices.values.items():
            if text in (label.casefold(), value.casefold()):
                return value
        raise ValueError(f"valor no admitido «{raw}»; elige uno de la lista desplegable")
    if column.kind == "text":
        return str(raw).strip()
    if column.kind == "date":
        if isinstance(raw, datetime):
            return raw.date()
        if isinstance(raw, date):
            return raw
        return _parse_date(str(raw))
    if isinstance(raw, bool):
        raise ValueError(f"se esperaba un número y hay «{raw}»")
    try:
        number = float(raw) if isinstance(raw, int | float) else parse_number(str(raw))
    except ValueError:
        raise ValueError(f"se esperaba un número y hay «{raw}»") from None
    if column.kind == "integer":
        if not number.is_integer():
            raise ValueError(f"se esperaba un número entero y hay «{raw}»")
        return int(number)
    return number


_PYDANTIC_MESSAGES = {
    "missing": "falta el dato",
    "greater_than": "debe ser mayor que {gt}",
    "greater_than_equal": "debe ser mayor o igual que {ge}",
    "less_than": "debe ser menor que {lt}",
    "less_than_equal": "debe ser menor o igual que {le}",
    "string_pattern_mismatch": "formato no válido",
    "string_too_short": "no puede estar vacío",
    "finite_number": "debe ser un número finito",
}


def _spanish(error: Mapping[str, object]) -> str:
    """Pydantic's message in Spanish (the project's own validators already speak Spanish)."""
    kind = str(error.get("type", ""))
    message = str(error.get("msg", ""))
    if kind == "value_error":
        return message.removeprefix("Value error, ")
    template = _PYDANTIC_MESSAGES.get(kind)
    if template is None:
        return message
    context = error.get("ctx")
    return template.format(**context) if isinstance(context, dict) else template


def _header_map(worksheet: Worksheet) -> dict[str, int]:
    first = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    return {
        str(value).strip().casefold(): index
        for index, value in enumerate(first)
        if value is not None and str(value).strip()
    }


def _rows(worksheet: Worksheet) -> list[tuple[int, tuple[object, ...]]]:
    return [
        (number, row)
        for number, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2)
        if any(cell is not None and str(cell).strip() for cell in row)
    ]


def _read_columns(
    sheet_name: str,
    columns: Sequence[Column],
    headers: Mapping[str, int],
    number: int,
    row: tuple[object, ...],
    errors: list[str],
) -> dict[str, object]:
    values: dict[str, object] = {}
    for column in columns:
        index = headers.get(column.header.casefold())
        raw = row[index] if index is not None and index < len(row) else None
        try:
            values[column.field] = _convert(raw, column)
        except ValueError as error:
            errors.append(f"Hoja «{sheet_name}», fila {number}, columna «{column.header}»: {error}")
    return values


def _validate(
    model: type[BaseModel],
    values: Mapping[str, object],
    where: str,
    headers_by_field: Mapping[str, str],
    errors: list[str],
) -> BaseModel | None:
    try:
        return model.model_validate(values)
    except ValidationError as error:
        for detail in error.errors():
            location = detail.get("loc", ())
            column = headers_by_field.get(str(location[0])) if location else None
            prefix = f"{where}, columna «{column}»" if column else where
            errors.append(f"{prefix}: {_spanish(detail)}")
        return None


def _missing_columns(sheet: Sheet, headers: Mapping[str, int], errors: list[str]) -> bool:
    missing = [c.header for c in sheet.columns if c.header.casefold() not in headers]
    if missing:
        errors.append(f"Hoja «{sheet.name}»: faltan las columnas {', '.join(missing)}")
    return bool(missing)


def _required_present(
    columns: Sequence[Column], values: Mapping[str, object], where: str, errors: list[str]
) -> bool:
    absent = [c.header for c in columns if c.required and values.get(c.field) is None]
    if absent:
        errors.append(f"{where}: faltan datos obligatorios en {', '.join(absent)}")
    return not absent


def _read_sheet(sheet: Sheet, worksheet: Worksheet, errors: list[str]) -> list[BaseModel]:
    headers = _header_map(worksheet)
    if _missing_columns(sheet, headers, errors):
        return []
    by_field = {column.field: column.header for column in sheet.columns}
    records = []
    for number, row in _rows(worksheet):
        where = f"Hoja «{sheet.name}», fila {number}"
        values = _read_columns(sheet.name, sheet.columns, headers, number, row, errors)
        if not _required_present(sheet.columns, values, where, errors):
            continue
        record = _validate(sheet.model, values, where, by_field, errors)
        if record is not None:
            records.append(record)
    return records


def _read_financials(
    worksheet: Worksheet, errors: list[str], warnings: list[str]
) -> list[BaseModel]:
    """Wide layout (one column per line item) turned into one fact per non-empty cell."""
    headers = _header_map(worksheet)
    if _missing_columns(FINANCIALS, headers, errors):
        return []
    facts: list[BaseModel] = []
    for number, row in _rows(worksheet):
        where = f"Hoja «{FINANCIALS.name}», fila {number}"
        keys = _read_columns(FINANCIALS.name, FINANCIALS.columns, headers, number, row, errors)
        if not _required_present(FINANCIALS.columns, keys, where, errors):
            continue
        if keys["period_end"] == keys["published_at"]:
            warnings.append(
                f"{where}: la fecha de publicación coincide con la de cierre. Si no es la fecha "
                "real de publicación, los datos se darán por conocidos antes de tiempo."
            )
        items = _read_columns(FINANCIALS.name, ITEM_COLUMNS, headers, number, row, errors)
        for column in ITEM_COLUMNS:
            value = items.get(column.field)
            if value is None:
                continue
            item = FinancialItem(column.field)
            fact_values = {
                **keys,
                "item": item.value,
                "value": value,
                "currency": keys["currency"] if item.is_monetary else None,
            }
            fact = _validate(
                FinancialFact, fact_values, f"{where}, columna «{column.header}»", {}, errors
            )
            if fact is not None:
                facts.append(fact)
    return facts


def _limit(errors: list[str]) -> list[str]:
    if len(errors) <= MAX_REPORTED_ERRORS:
        return errors
    return [*errors[:MAX_REPORTED_ERRORS], f"… y {len(errors) - MAX_REPORTED_ERRORS} más"]


def read_workbook(path: Path) -> WorkbookImport:
    """Validate and read a workbook. Raises ``WorkbookError`` listing every problem found."""
    if not path.is_file():
        raise WorkbookError([f"No se encuentra el fichero {path}"])
    try:
        workbook = load_workbook(path, data_only=True)
    except Exception as error:  # openpyxl raises several kinds of errors on unreadable files
        raise WorkbookError(
            [f"No se puede abrir {path.name} como libro de Excel: {error}"]
        ) from error
    errors: list[str] = []
    warnings: list[str] = []
    collections: dict[str, list[BaseModel]] = {}
    for sheet in SHEETS:
        worksheet = workbook[sheet.name] if sheet.name in workbook.sheetnames else None
        if not isinstance(worksheet, Worksheet):
            if sheet.required:
                errors.append(f"Falta la hoja «{sheet.name}»")
            continue
        if sheet is FINANCIALS:
            records = _read_financials(worksheet, errors, warnings)
        else:
            records = _read_sheet(sheet, worksheet, errors)
        if sheet.required and not records and not errors:
            errors.append(f"La hoja «{sheet.name}» está vacía")
        collections[sheet.collection] = records
    if errors:
        raise WorkbookError(_limit(errors))
    try:
        bundle = DataBundle.model_validate(collections)
    except ValidationError as error:
        raise WorkbookError(_limit([_spanish(detail) for detail in error.errors()])) from error
    return WorkbookImport(bundle, tuple(warnings))


class WorkbookProvider(DataProvider):
    """Provider backed by an Excel workbook filled by the analysts."""

    name: ClassVar[str] = "excel"

    def __init__(self, path: Path) -> None:
        self.path = path

    def fetch(self) -> DataBundle:
        return read_workbook(self.path).bundle


# --- Writing (template and example) ----------------------------------------------------------
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="1F5F8B")
_TITLE_FONT = Font(bold=True, size=14)
_BOLD = Font(bold=True)
_DATE_FORMAT = "DD/MM/YYYY"

_INSTRUCTIONS = (
    "1. Rellena las hojas con datos de Bloomberg o FactSet (por ejemplo, con sus complementos de "
    "Excel) o a mano. Si usas fórmulas, guarda el libro después de que se calculen.",
    "2. Cada libro es una carga completa: incluye siempre las empresas (Emisores), sus acciones y "
    "sus bonos, aunque solo quieras actualizar precios.",
    "3. Una celda vacía significa «sin dato» y la herramienta nunca la rellena. No pongas cero si "
    "no lo es.",
    "4. Importes en millones de la divisa de reporte; gastos financieros y capex en positivo; "
    "fechas como fechas de Excel (dd/mm/aaaa).",
    "5. La «Fecha de publicación» de los estados financieros es imprescindible: impide usar "
    "cifras antes de que se conocieran.",
    "6. Carga el libro con: .\\analisis importar RUTA_DEL_LIBRO.xlsx. Si hay errores, la "
    "herramienta indica hoja, fila y columna, y no carga nada.",
    "7. Guarda los libros con datos reales en data\\import o fuera del repositorio (git ignora "
    "esa carpeta): sus licencias no permiten publicarlos.",
    "Los códigos de Bloomberg y FactSet son orientativos: compruébalos en el terminal (FLDS) o en "
    "la ayuda de FactSet antes de usarlos, y confirma que el contrato permite este uso de los "
    "datos.",
)


def _display(value: object, column: Column) -> object:
    """Model value as written in the workbook: the label for a choice, the value otherwise."""
    if value is None or column.choices is None:
        return value
    stored = getattr(value, "value", value)
    return next((label for label, v in column.choices.values.items() if v == stored), stored)


def _financial_rows(bundle: DataBundle) -> list[list[object]]:
    grouped: dict[tuple[str, date, str, date], dict[FinancialItem, FinancialFact]] = defaultdict(
        dict
    )
    for fact in bundle.financial_facts:
        key = (fact.issuer_id, fact.period_end, fact.fiscal_period.value, fact.published_at)
        grouped[key][fact.item] = fact
    rows: list[list[object]] = []
    for (issuer_id, period_end, period, published_at), facts in sorted(grouped.items()):
        currency = next((fact.currency for fact in facts.values() if fact.currency), None)
        values = [facts[item].value if item in facts else None for item in FinancialItem]
        rows.append([issuer_id, period, period_end, published_at, currency, *values])
    return rows


def _list_ranges() -> dict[str, str]:
    """Cell range of each list in the hidden sheet (one column per list)."""
    return {
        choices.name: (
            f"{LISTS_SHEET}!${get_column_letter(index)}$2:"
            f"${get_column_letter(index)}${len(choices.values) + 1}"
        )
        for index, choices in enumerate(ALL_CHOICES, start=1)
    }


def _write_lists(workbook: Workbook) -> None:
    lists = workbook.create_sheet(LISTS_SHEET)
    for index, choices in enumerate(ALL_CHOICES, start=1):
        lists.cell(row=1, column=index, value=choices.name).font = _BOLD
        for row, label in enumerate(choices.values, start=2):
            lists.cell(row=row, column=index, value=label)
    lists.sheet_state = "hidden"


def _write_data_sheet(
    worksheet: Worksheet, columns: Sequence[Column], rows: Sequence[Sequence[object]]
) -> None:
    worksheet.append([column.header for column in columns])
    for cell in worksheet[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    worksheet.freeze_panes = "A2"
    for row in rows:
        worksheet.append(list(row))
    ranges = _list_ranges()
    for index, column in enumerate(columns, start=1):
        letter = get_column_letter(index)
        worksheet.column_dimensions[letter].width = max(12, len(column.header) + 4)
        if column.kind == "date":
            for (cell,) in worksheet.iter_rows(min_row=2, min_col=index, max_col=index):
                cell.number_format = _DATE_FORMAT
        if column.choices is not None:
            validation = DataValidation(
                type="list",
                formula1=f"={ranges[column.choices.name]}",
                allow_blank=True,
                showErrorMessage=True,
                errorTitle="Valor no admitido",
                error="Elige un valor de la lista desplegable.",
            )
            validation.add(f"{letter}2:{letter}{LAST_VALIDATED_ROW}")
            worksheet.add_data_validation(validation)


def _write_instructions(worksheet: Worksheet) -> None:
    worksheet.title = "Instrucciones"
    worksheet.append(["Plantilla de importación de datos reales"])
    worksheet["A1"].font = _TITLE_FONT
    for line in _INSTRUCTIONS:
        worksheet.append([line])
    for sheet in SHEETS:
        worksheet.append([])
        worksheet.append([f"Hoja «{sheet.name}»", sheet.description])
        worksheet.cell(row=worksheet.max_row, column=1).font = _BOLD
        worksheet.append(
            [
                "Columna",
                "Obligatoria",
                "Qué poner",
                "Bloomberg (orientativo)",
                "FactSet (orientativo)",
            ]
        )
        for cell in worksheet[worksheet.max_row]:
            cell.font = _HEADER_FONT
            cell.fill = _HEADER_FILL
        for column in sheet_columns(sheet):
            required = "Sí" if column.required else "No"
            worksheet.append(
                [column.header, required, column.help, column.bloomberg, column.factset]
            )
    for letter, width in zip("ABCDE", (34, 12, 70, 36, 30), strict=True):
        worksheet.column_dimensions[letter].width = width
    for row in worksheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def write_workbook(bundle: DataBundle, path: Path) -> Path:
    """Write the template; with a non-empty bundle, filled with its data (e.g. the sample)."""
    workbook = Workbook()
    instructions = workbook.active
    assert isinstance(instructions, Worksheet)
    _write_instructions(instructions)
    for sheet in SHEETS:
        if sheet is FINANCIALS:
            rows = _financial_rows(bundle)
        else:
            rows = [
                [_display(getattr(record, column.field), column) for column in sheet.columns]
                for record in getattr(bundle, sheet.collection)
            ]
        _write_data_sheet(workbook.create_sheet(sheet.name), sheet_columns(sheet), rows)
    _write_lists(workbook)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path
